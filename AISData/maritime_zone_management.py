"""Validate editable maritime areas and tabular imports before any write."""

import ast
import csv
import io
import json
import re
import zipfile
from pathlib import Path

from django.db import transaction

from .maritime_zone_registry import (
    CUSTOM_ID_OFFSET, get_zone_snapshot, invalidate_zone_snapshot, merge_zones,
)
from .maritime_zones import MaritimeZoneDataError, SUPPORTED_ZONE_TYPES
from .models import MaritimeZoneOverride, MaritimeZoneRevision
from .monitor_area_geometry import normalise_polygon_vertices


MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
TYPE_LABELS = {
    "PRT": "港口", "ANC": "锚地", "BTH": "泊位", "LIQ": "液体散货码头",
    "GCO": "杂货码头", "DRY": "干散货码头", "CTR": "集装箱码头",
    "GAS": "气体码头", "ROR": "滚装码头", "PAX": "客运码头",
}
HEADERS = ["id", "AOI_Name", "type", "locode", "AOI_Description", "geometry", "is_active"]
ALIASES = {
    "id": "id", "编号": "id", "区域编号": "id",
    "aoi_name": "name", "name": "name", "名称": "name", "区域名称": "name",
    "type": "zone_type", "zone_type": "zone_type", "类型": "zone_type",
    "locode": "locode", "港口代码": "locode",
    "aoi_description": "description", "description": "description", "说明": "description",
    "geometry": "vertices", "vertices": "vertices", "坐标": "vertices", "顶点坐标": "vertices",
    "is_active": "is_active", "启用": "is_active", "是否启用": "is_active",
}


class ZoneConflict(ValueError):
    pass


def serialize_zone(zone):
    return {
        "id": zone.source_id, "name": zone.name, "description": zone.description,
        "locode": zone.locode, "zone_type": zone.zone_type,
        "vertices": [list(point) for point in zone.points],
        "is_active": zone.is_active, "source": zone.source,
        "type_label": TYPE_LABELS[zone.zone_type],
    }


def parse_vertices(value):
    if isinstance(value, str):
        if len(value) > 100000:
            raise ValueError("坐标内容过长")
        value = value.strip()
        if value.upper().startswith("POLYGON"):
            from shapely import wkt
            try:
                polygon = wkt.loads(value)
            except Exception as exc:
                raise ValueError("WKT 多边形格式不正确") from exc
            if polygon.geom_type != "Polygon" or polygon.interiors or polygon.has_z or polygon.is_empty:
                raise ValueError("仅支持不带内环的二维多边形")
            value = [list(point) for point in polygon.exterior.coords]
        elif value.startswith(("[", "{", "(")):
            try:
                value = json.loads(value)
            except ValueError:
                try:
                    value = ast.literal_eval(value)
                except (SyntaxError, ValueError, RecursionError) as exc:
                    raise ValueError("坐标数组格式不正确") from exc
            if isinstance(value, tuple):
                value = list(value)
        else:
            value = [re.split(r"[\s,，]+", row.strip())
                     for row in re.split(r"[;；\n]+", value) if row.strip()]
    if isinstance(value, dict):
        rings = value.get("coordinates")
        if value.get("type") != "Polygon" or not isinstance(rings, list) or len(rings) != 1:
            raise ValueError("仅支持不带内环的 GeoJSON Polygon")
        value = rings[0]
    if isinstance(value, list) and any(
        isinstance(number, bool) for point in value if isinstance(point, (list, tuple)) for number in point
    ):
        raise ValueError("经纬度必须是数字")
    return normalise_polygon_vertices(value, label="海事区域", max_vertices=500)


def normalize_row(data):
    if not isinstance(data, dict):
        raise ValueError("每条区域必须是对象")
    row = {}
    for key, value in data.items():
        canonical = ALIASES.get(str(key).strip().lower())
        if canonical:
            if canonical in row:
                raise ValueError(f"重复字段：{canonical}")
            row[canonical] = value
    name = str(row.get("name") or "").strip()
    if not name or len(name) > 255:
        raise ValueError("名称必填，最多 255 个字符")
    zone_type = str(row.get("zone_type") or "").strip()
    zone_type = {label: code for code, label in TYPE_LABELS.items()}.get(zone_type, zone_type.upper())
    if zone_type not in SUPPORTED_ZONE_TYPES:
        raise ValueError("类型须为港口/PRT、锚地/ANC、泊位/BTH 或模板中的码头类型")
    locode = str(row.get("locode") or "CN").strip().upper()
    if not re.fullmatch(r"[A-Z]{2}[A-Z0-9]{0,6}", locode):
        raise ValueError("LOCODE 须为 2～8 位字母数字，前两位为国家/地区字母代码")
    description = str(row.get("description") or "").strip()
    if len(description) > 2000:
        raise ValueError("说明最多 2000 个字符")
    raw_active = row.get("is_active", True)
    active_text = str(raw_active).strip().lower()
    if raw_active is None or active_text == "":
        active = True
    elif active_text in {"true", "1", "1.0", "是", "启用"}:
        active = True
    elif active_text in {"false", "0", "0.0", "否", "停用"}:
        active = False
    else:
        raise ValueError("启用状态须为 true/false、1/0 或 是/否")
    identifier = row.get("id")
    if identifier is not None and str(identifier).strip():
        text = str(identifier).strip()
        if not re.fullmatch(r"-?[0-9]+(?:\.0+)?", text):
            raise ValueError("编号须为整数；新增区域请留空")
        identifier = int(text.split(".")[0])
        if identifier == 0 or abs(identifier) > 9_007_199_254_740_991:
            raise ValueError("编号超出有效范围")
    else:
        identifier = None
    return {
        "id": identifier, "name": name, "zone_type": zone_type,
        "locode": locode, "description": description,
        "vertices": parse_vertices(row.get("vertices")), "is_active": active,
    }


def validate_rows(rows, snapshot=None):
    if not isinstance(rows, list) or not 1 <= len(rows) <= MAX_IMPORT_ROWS:
        raise ValueError(f"每次须提供 1～{MAX_IMPORT_ROWS} 条区域")
    snapshot = snapshot or get_zone_snapshot()
    existing = {zone.source_id for zone in snapshot["zones"]}
    cleaned, seen = [], set()
    for index, data in enumerate(rows, start=2):
        try:
            row = normalize_row(data)
            identifier = row["id"]
            if identifier is not None:
                if identifier not in existing:
                    raise ValueError(f"编号 {identifier} 不存在；新增区域请将编号留空")
                if identifier in seen:
                    raise ValueError(f"编号 {identifier} 在文件中重复")
                seen.add(identifier)
            cleaned.append(row)
        except (ValueError, TypeError) as exc:
            raise ValueError(f"第 {index} 行：{exc}") from exc
    return cleaned


def read_import(upload):
    if upload is None:
        raise ValueError("请选择 CSV 或 Excel 文件")
    if upload.size > MAX_IMPORT_BYTES:
        raise ValueError("文件不能超过 5 MB")
    suffix = Path(upload.name).suffix.lower()
    raw = upload.read(MAX_IMPORT_BYTES + 1)
    if not raw or len(raw) > MAX_IMPORT_BYTES:
        raise ValueError("文件为空或超过 5 MB")
    workbook = None
    try:
        if suffix == ".csv":
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                try:
                    text = raw.decode("gb18030")
                except UnicodeDecodeError as exc:
                    raise ValueError("CSV 请使用 UTF-8 或 GB18030 编码") from exc
            rows = csv.reader(io.StringIO(text, newline=""), strict=True)
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError("Excel 解压内容超过 30 MB，请拆分文件")
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            # The first sheet is the data sheet; formulas are never evaluated.
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()
            def excel_rows():
                for number, cells in enumerate(sheet.iter_rows(), start=1):
                    if any(cell.data_type in {"f", "e"} for cell in cells):
                        raise ValueError(f"第 {number} 行含公式或错误值，请粘贴为数值后导入")
                    yield [cell.value for cell in cells]
            rows = excel_rows()
        elif suffix == ".xls":
            import xlrd
            workbook = xlrd.open_workbook(file_contents=raw, on_demand=True)
            sheet = workbook.sheet_by_index(0)
            rows = (sheet.row_values(number) for number in range(sheet.nrows))
        else:
            raise ValueError("仅支持 .csv、.xlsx、.xls 文件")
        headers = next(rows, None)
        if not headers or len(headers) > 32:
            raise ValueError("文件缺少表头或列数超过 32")
        keys = [ALIASES.get(str(value or "").strip().lower()) for value in headers]
        recognized = [key for key in keys if key]
        if len(recognized) != len(set(recognized)):
            raise ValueError("表头包含重复字段")
        if not {"name", "zone_type", "vertices"}.issubset(recognized):
            raise ValueError("表头至少须包含 AOI_Name/名称、type/类型、geometry/坐标")
        result = []
        for number, values in enumerate(rows, start=2):
            if number > MAX_IMPORT_ROWS + 1:
                raise ValueError(f"最多导入 {MAX_IMPORT_ROWS} 行，请删除尾部空行或拆分文件")
            if not any(value is not None and str(value).strip() for value in values):
                # Keep row numbers exact without silently accepting blank data rows.
                result.append(None)
                continue
            if len(values) > len(headers) and any(str(v or "").strip() for v in values[len(headers):]):
                raise ValueError(f"第 {number} 行列数超过表头；CSV 坐标须放在带引号的单元格中")
            result.append({key: values[index] if index < len(values) else None
                           for index, key in enumerate(keys) if key})
        while result and result[-1] is None:
            result.pop()
        if None in result:
            raise ValueError(f"第 {result.index(None) + 2} 行为空，请删除空行")
        return validate_rows(result)
    except ImportError as exc:
        raise ValueError("Excel 读取依赖未安装，请安装 requirements.txt 中的 openpyxl、xlrd") from exc
    except (ValueError, csv.Error, MaritimeZoneDataError):
        raise
    except Exception as exc:
        raise ValueError("文件无法读取，请确认文件未损坏且格式与扩展名一致") from exc
    finally:
        if workbook is not None:
            if suffix == ".xls":
                workbook.release_resources()
            else:
                workbook.close()


def save_rows(rows, expected_revision):
    if isinstance(expected_revision, bool) or not isinstance(expected_revision, int):
        raise ValueError("缺少数据版本，请刷新列表或重新预览文件")
    with transaction.atomic():
        state, _ = MaritimeZoneRevision.objects.select_for_update().get_or_create(pk=1)
        if state.version != expected_revision:
            raise ZoneConflict("区域数据已被更新，请刷新列表或重新预览后再保存")
        records = list(MaritimeZoneOverride.objects.all())
        snapshot = {"revision": state.version, "zones": merge_zones(records)}
        cleaned = validate_rows(rows, snapshot)
        saved = []
        for row in cleaned:
            identifier = row.pop("id")
            if identifier is None:
                record = MaritimeZoneOverride.objects.create(**row)
            elif identifier >= CUSTOM_ID_OFFSET:
                record = MaritimeZoneOverride.objects.get(pk=identifier - CUSTOM_ID_OFFSET, source_id=None)
                for field, value in row.items():
                    setattr(record, field, value)
                record.save()
            else:
                record, _ = MaritimeZoneOverride.objects.update_or_create(source_id=identifier, defaults=row)
            saved.append(record.source_id if record.source_id is not None else CUSTOM_ID_OFFSET + record.pk)
        state.version += 1
        state.save(update_fields=["version"])
        transaction.on_commit(invalidate_zone_snapshot)
    invalidate_zone_snapshot()
    return {"ids": saved, "revision": state.version, "count": len(saved)}


def table_rows(zones):
    return [[zone.source_id, zone.name, zone.zone_type, zone.locode, zone.description,
             json.dumps([list(point) for point in zone.points], separators=(",", ":")),
             "true" if zone.is_active else "false"] for zone in zones]
