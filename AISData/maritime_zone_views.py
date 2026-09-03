import csv
import io
import json

from django.http import HttpResponse, JsonResponse
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET, require_POST, require_http_methods

from .maritime_zone_management import (
    HEADERS, TYPE_LABELS, ZoneConflict, read_import, save_rows, serialize_zone, table_rows,
)
from .maritime_zone_registry import get_zone_snapshot
from .maritime_zones import MaritimeZoneDataError


def error_response(exc):
    if isinstance(exc, MaritimeZoneDataError):
        return JsonResponse({"success": False, "message": "海事区域数据暂时不可用"}, status=503)
    return JsonResponse({"success": False, "message": str(exc)}, status=409 if isinstance(exc, ZoneConflict) else 400)


def read_json(request):
    try:
        data = json.loads(request.body)
    except (ValueError, UnicodeDecodeError) as exc:
        raise ValueError("请求体必须是有效的 JSON 对象") from exc
    if not isinstance(data, dict):
        raise ValueError("请求体必须是 JSON 对象")
    return data


@never_cache
@require_http_methods(["GET", "POST"])
def zone_management(request):
    try:
        if request.method == "GET":
            snapshot = get_zone_snapshot()
            return JsonResponse({
                "success": True, "revision": snapshot["revision"],
                "zones": [serialize_zone(zone) for zone in snapshot["zones"]],
                "types": TYPE_LABELS,
            })
        data = read_json(request)
        result = save_rows(data.get("rows"), data.get("revision"))
        return JsonResponse({"success": True, "message": f"已保存 {result['count']} 条区域，下一轮检测生效", **result})
    except (ValueError, MaritimeZoneDataError) as exc:
        return error_response(exc)


@never_cache
@require_POST
def zone_import_preview(request):
    try:
        rows = read_import(request.FILES.get("file"))
        created = sum(row["id"] is None for row in rows)
        return JsonResponse({
            "success": True, "revision": get_zone_snapshot()["revision"],
            "rows": rows, "created": created, "updated": len(rows) - created,
        })
    except (ValueError, csv.Error, MaritimeZoneDataError) as exc:
        return error_response(exc)


@never_cache
@require_GET
def zone_export(request):
    kind = request.GET.get("format", "csv").lower()
    if kind not in {"csv", "xlsx"}:
        return error_response(ValueError("导出格式须为 csv 或 xlsx"))
    is_template = request.GET.get("template") == "1"
    try:
        rows = [["", "示例锚地（请修改）", "ANC", "CN", "填写区域说明",
                 "[[113.8,22.5],[113.81,22.5],[113.81,22.51],[113.8,22.51]]", "true"]] if is_template else table_rows(get_zone_snapshot()["zones"])
        if kind == "csv":
            output = io.StringIO(newline="")
            writer = csv.writer(output)
            writer.writerow(HEADERS)
            # Quoting alone does not prevent spreadsheet formula execution.
            writer.writerows([
                ["'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")) else value for value in row]
                for row in rows
            ])
            response = HttpResponse(output.getvalue().encode("utf-8-sig"), content_type="text/csv; charset=utf-8")
        else:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "区域数据"
            sheet.append(HEADERS)
            for row in rows:
                sheet.append(row)
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="164E63")
            for column, width in zip("ABCDEFG", [18, 32, 12, 14, 38, 75, 12]):
                sheet.column_dimensions[column].width = width
            help_sheet = workbook.create_sheet("填写说明")
            for text in [
                "每行一条区域；只读取第一个工作表。", "编号留空为新增，保留编号为更新该区域。",
                "必填列：AOI_Name、type、geometry。", "type：" + "；".join(f"{k}={v}" for k, v in TYPE_LABELS.items()),
                "geometry：WGS84 经纬度数组 [[经度,纬度],...]，至少 3 点；支持单环 WKT Polygon。",
                "is_active：true 启用，false 停用；空值默认为启用。", "先上传预览，确认后整批保存；有错误时整批不保存。",
                "最多 2000 行、每区 500 点、文件 5 MB；请勿使用公式。",
            ]:
                help_sheet.append([text])
            help_sheet.column_dimensions["A"].width = 115
            output = io.BytesIO()
            workbook.save(output)
            workbook.close()
            response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        filename = "maritime-zone-template" if is_template else "maritime-zones"
        response["Content-Disposition"] = f'attachment; filename="{filename}.{kind}"'
        return response
    except MaritimeZoneDataError as exc:
        return error_response(exc)
