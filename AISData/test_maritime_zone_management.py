import csv
import io
import json
from unittest.mock import patch

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from openpyxl import Workbook, load_workbook

from .maritime_zone_management import HEADERS, ZoneConflict, save_rows, serialize_zone, validate_rows
from .maritime_zone_registry import built_in_zones, get_zone_snapshot, maritime_zone_snapshot
from .maritime_zones import PORT_ZONE_TYPES, find_authorized_anchorage, get_maritime_zones, zones_containing_point
from .models import MaritimeZoneOverride, MaritimeZoneRevision


MANAGE = "/AISData/maritime-zones/manage/"
PREVIEW = "/AISData/maritime-zones/import-preview/"
EXPORT = "/AISData/maritime-zones/export/"
POINTS = [[114.60, 22.50], [114.61, 22.50], [114.61, 22.51], [114.60, 22.51]]


@override_settings(CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}})
class MaritimeZoneManagementTests(TestCase):
    def setUp(self):
        MaritimeZoneRevision.objects.get_or_create(pk=1)
        cache.clear()

    def row(self, **changes):
        return {"name": "自定义区域", "zone_type": "ANC", "vertices": POINTS, **changes}

    def commit(self, rows, revision=None):
        if revision is None:
            revision = get_zone_snapshot()["revision"]
        return self.client.post(MANAGE, data=json.dumps({"revision": revision, "rows": rows}), content_type="application/json")

    def csv_file(self, rows, headers=HEADERS, encoding="utf-8-sig"):
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)
        return SimpleUploadedFile("areas.csv", output.getvalue().encode(encoding))

    def test_custom_anchorage_is_used_by_detection_and_geojson(self):
        from IllegalAnchored.zones import classify_location
        self.assertEqual(classify_location(114.605, 22.505)["state"], "unassigned")
        response = self.commit([self.row(name="新增锚地")])
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(classify_location(114.605, 22.505)["zone_name"], "新增锚地")
        payload = self.client.get("/AISData/maritime-zones/?type=ANC").json()
        self.assertTrue(any(feature["properties"]["name"] == "新增锚地" for feature in payload["features"]))

    def test_new_berth_is_normal_port_operation_and_legal_berthing(self):
        from IllegalAnchored.views import _is_normal_port_operation
        from AbnormalParking.utils import _legal_berthing, _near_fixed_facility, get_parking_config
        response = self.commit([self.row(zone_type="BTH", name="测试泊位")])
        self.assertEqual(response.status_code, 200, response.content)
        ship = {"lon": 114.605, "lat": 22.505, "matched_port_name": ""}
        self.assertTrue(zones_containing_point(ship["lon"], ship["lat"], PORT_ZONE_TYPES))
        self.assertTrue(_is_normal_port_operation(ship, {"state": "unassigned"}))
        self.assertFalse(_is_normal_port_operation(ship, {"state": "prohibited"}))
        self.assertEqual(_legal_berthing(ship, get_parking_config()), (True, "测试泊位"))
        self.assertTrue(_near_fixed_facility(ship, get_parking_config())[3])

    def test_move_and_disable_existing_area_preserve_identifier(self):
        created = self.commit([self.row()]).json()
        identifier = created["ids"][0]
        moved = [[lon + .1, lat] for lon, lat in POINTS]
        self.assertEqual(self.commit([self.row(id=identifier, vertices=moved)]).status_code, 200)
        self.assertIsNone(find_authorized_anchorage(114.605, 22.505))
        self.assertEqual(find_authorized_anchorage(114.705, 22.505).source_id, identifier)
        self.assertEqual(self.commit([self.row(id=identifier, vertices=moved, is_active=False)]).status_code, 200)
        self.assertNotIn(identifier, [zone.source_id for zone in zones_containing_point(114.705, 22.505, {"ANC"})])
        self.assertEqual(MaritimeZoneOverride.objects.count(), 1)

    def test_published_area_update_cannot_fall_back_to_old_coordinates(self):
        from IllegalAnchored.zones import classify_location
        base = next(zone for zone in built_in_zones() if zone.name == "大鹏湾LNG专用锚地")
        lon = sum(point[0] for point in base.points) / len(base.points)
        lat = sum(point[1] for point in base.points) / len(base.points)
        with patch("AISData.maritime_zone_registry.built_in_zones", return_value=(base,)):
            row = serialize_zone(base)
            row["vertices"] = POINTS
            self.assertEqual(self.commit([row]).status_code, 200)
            self.assertNotEqual(classify_location(lon, lat)["state"], "authorized")
            self.assertEqual(classify_location(114.605, 22.505)["zone_name"], base.name)

    def test_bundled_port_edit_overrides_original_polygon(self):
        base = next(zone for zone in built_in_zones() if zone.zone_type == "PRT")
        with patch("AISData.maritime_zone_registry.built_in_zones", return_value=(base,)):
            row = serialize_zone(base)
            row.update(vertices=POINTS, name="更新港口")
            self.assertEqual(self.commit([row]).status_code, 200)
            zones = get_maritime_zones({"PRT"})
            self.assertEqual(len(zones), 1)
            self.assertEqual(zones[0].source_id, base.source_id)
            self.assertEqual(zones[0].points, tuple(map(tuple, POINTS)))

    def test_preview_is_read_only_and_reports_create_update(self):
        base = next(zone for zone in built_in_zones() if zone.zone_type == "PRT")
        upload = self.csv_file([
            [base.source_id, "更新港口", "PRT", "CNSZX", "", json.dumps(POINTS), "true"],
            ["", "新增泊位", "泊位", "CN", "", json.dumps(POINTS), "是"],
        ])
        response = self.client.post(PREVIEW, {"file": upload})
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual((response.json()["created"], response.json()["updated"]), (1, 1))
        self.assertEqual(MaritimeZoneOverride.objects.count(), 0)
        result = self.commit(response.json()["rows"], response.json()["revision"])
        self.assertEqual(result.status_code, 200, result.content)
        self.assertEqual(MaritimeZoneOverride.objects.count(), 2)

    def test_chinese_gb18030_csv_and_wkt(self):
        upload = self.csv_file([["中文锚地", "锚地", "POLYGON ((114.60 22.50,114.61 22.50,114.61 22.51,114.60 22.50))"]],
                               headers=["名称", "类型", "坐标"], encoding="gb18030")
        response = self.client.post(PREVIEW, {"file": upload})
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["rows"][0]["zone_type"], "ANC")

    def test_xlsx_preview_and_formula_rejection(self):
        for name, status in [("测试泊位", 200), ("=1+1", 400)]:
            workbook = Workbook()
            workbook.active.append(HEADERS)
            workbook.active.append([None, name, "BTH", "CN", "", json.dumps(POINTS), True])
            output = io.BytesIO()
            workbook.save(output)
            response = self.client.post(PREVIEW, {"file": SimpleUploadedFile("areas.xlsx", output.getvalue())})
            self.assertEqual(response.status_code, status, response.content)

    def test_export_round_trips_all_current_areas(self):
        for kind in ["csv", "xlsx"]:
            response = self.client.get(EXPORT, {"format": kind})
            self.assertEqual(response.status_code, 200)
            preview = self.client.post(PREVIEW, {"file": SimpleUploadedFile(f"areas.{kind}", response.content)})
            self.assertEqual(preview.status_code, 200, preview.content)
            self.assertEqual(preview.json()["created"], 0)
            self.assertEqual(preview.json()["updated"], len(built_in_zones()))

    def test_export_protects_formula_strings(self):
        self.commit([self.row(name="=1+1")])
        response = self.client.get(EXPORT, {"format": "xlsx"})
        workbook = load_workbook(io.BytesIO(response.content), data_only=False)
        cells = [cell for row in workbook.active for cell in row if cell.value == "=1+1"]
        self.assertEqual(len(cells), 1)
        self.assertEqual(cells[0].data_type, "s")
        response = self.client.get(EXPORT, {"format": "csv"})
        self.assertIn("'=1+1", response.content.decode("utf-8-sig"))

    def test_bad_import_reports_line_and_writes_nothing(self):
        upload = self.csv_file([
            ["", "有效", "ANC", "CN", "", json.dumps(POINTS), "true"],
            ["", "无效", "ANC", "CN", "", "[[181,22],[182,22],[181,23]]", "true"],
        ])
        response = self.client.post(PREVIEW, {"file": upload})
        self.assertEqual(response.status_code, 400)
        self.assertIn("第 3 行", response.json()["message"])
        self.assertEqual(MaritimeZoneOverride.objects.count(), 0)

    def test_commit_validates_entire_batch_before_writing(self):
        response = self.commit([self.row(), self.row(vertices=[[1, 1], [2, 2], [3, 3]])])
        self.assertEqual(response.status_code, 400)
        self.assertEqual(MaritimeZoneOverride.objects.count(), 0)
        self.assertEqual(get_zone_snapshot()["revision"], 0)

    def test_invalid_geometry_and_duplicate_identifiers(self):
        identifier = self.commit([self.row()]).json()["ids"][0]
        for rows in [
            [self.row(vertices=[[0, 0], [1, 1], [0, 1], [1, 0]])],
            [self.row(vertices=[[1, 1], [2, float("inf")], [2, 3]])],
            [self.row(vertices={"type": "MultiPolygon", "coordinates": []})],
            [self.row(id=999999)],
            [self.row(id=identifier), self.row(id=identifier)],
        ]:
            self.assertEqual(self.commit(rows).status_code, 400)
        self.assertEqual(MaritimeZoneOverride.objects.count(), 1)

    def test_stale_editor_and_import_cannot_overwrite_newer_data(self):
        self.assertEqual(self.commit([self.row()], revision=0).status_code, 200)
        response = self.commit([self.row(name="过期编辑")], revision=0)
        self.assertEqual(response.status_code, 409)
        self.assertEqual(MaritimeZoneOverride.objects.count(), 1)

    def test_new_request_gets_update_and_one_snapshot_serves_repeated_queries(self):
        built_in_zones()
        with maritime_zone_snapshot():
            with self.assertNumQueries(3):
                get_maritime_zones()
                get_maritime_zones()
                get_maritime_zones({"ANC"})
        before = get_zone_snapshot()["revision"]
        self.commit([self.row()])
        with maritime_zone_snapshot():
            self.assertGreater(get_zone_snapshot()["revision"], before)
            self.assertIsNotNone(find_authorized_anchorage(114.605, 22.505))

    def test_shared_behavior_recomputes_when_zone_changes_for_same_ais_point(self):
        from .behavior_recognition import update_behavior_results
        ship = {"mmsi": "413000001", "timestamp": "2026-09-03T01:00:00+00:00",
                "lon": 114.605, "lat": 22.505, "speed": 0.1}
        with patch("AISData.behavior_recognition._calculate", return_value=({}, {})) as calculate:
            update_behavior_results([ship])
            update_behavior_results([ship])
            self.assertEqual(calculate.call_count, 1)
            self.commit([self.row()])
            update_behavior_results([ship])
            self.assertEqual(calculate.call_count, 2)

    def test_csrf_is_required_for_writes_and_preview(self):
        client = Client(enforce_csrf_checks=True)
        self.assertEqual(client.post(MANAGE, data="{}", content_type="application/json").status_code, 403)
        self.assertEqual(client.post(PREVIEW).status_code, 403)

    def test_empty_corrupt_and_unsupported_files_are_rejected(self):
        for name, content in [("bad.csv", b""), ("bad.xlsx", b"garbage"), ("bad.xls", b"garbage"), ("bad.exe", b"content")]:
            response = self.client.post(PREVIEW, {"file": SimpleUploadedFile(name, content)})
            self.assertEqual(response.status_code, 400, response.content)

    def test_page_contains_management_import_and_berth_layer(self):
        response = self.client.get("/")
        for identifier in ["btn-maritime", "maritime-import-file", "maritime-draw", "info-maritime-berths"]:
            self.assertContains(response, f'id="{identifier}"')
